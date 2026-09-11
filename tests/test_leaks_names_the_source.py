"""A worksheet row names the SOURCE DOCUMENT, on both passes.

An export is named for its source's SCRUBBED stem, deliberately: the `.txt` is
the artifact that gets shared, so a party or attorney name must not survive in
its filename. The consequence is that the export's name is a pseudonym —
`Feit Decl. ISO Mot..pdf` is written to `Kingscote Decl. ISO Mot..txt` — and
`--fix-leaks`, which works from the exports and never opens a PDF, wrote that
pseudonym into the File column. A delivered folder came back with three rows
naming a declaration the operator does not have ("There is no Kingscote
decl."), while the FULL run's rows named the source all along: one column, two
passes, two answers.

Run:  cd PDF-Linker && python3 -m pytest tests/test_leaks_names_the_source.py -v
"""
import logging
import pathlib
import sys
import types
import zipfile

import openpyxl

import pdf_linker as P

log = logging.getLogger("test")
BODY = ("Declaration of Ondine Feit in support of the motion.\n"
        "Respondent materially breached the Agreement as NAN alleged supra.\n")


def _docx(case, name, body_text=BODY):
    x = "".join(f"<w:p><w:r><w:t>{t}</w:t></w:r></w:p>"
                for t in body_text.split("\n"))
    with zipfile.ZipFile(case / name, "w") as z:
        z.writestr("word/document.xml",
                   '<?xml version="1.0"?><w:document xmlns:w="http://schemas.'
                   'openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                   + x + "</w:body></w:document>")


def _case(tmp_path, *names):
    case = tmp_path / "Case"
    case.mkdir()
    for n in names:
        _docx(case, n)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Case Number", "Title Plaintiff", "Title Defendant"])
    ws.append(["24STCV00001", "Ondine Feit", "Cascadia Freight, Inc."])
    wb.save(case / "Order_Mine.xlsx")
    return case


def _pz(*names):
    reg = P._PnFakeRegistry()
    return P.Pseudonymizer(P._pn_build_terms(list(names), [], [], registry=reg),
                           {}, registry=reg)


# ── the map ─────────────────────────────────────────────────────────────────

class TestTheExportSourceMap:
    def test_it_names_the_source_of_a_scrubbed_export(self, tmp_path):
        case = _case(tmp_path, "Feit Decl. ISO Mot..docx")
        pz = _pz("Ondine Feit")
        m = P._pn_export_sources(case, "Text Files", pz, log)
        # The export's own name is a pseudonym, and the map reads through it.
        assert list(m.values()) == ["Feit Decl. ISO Mot..docx"]
        (export,) = m
        assert "feit" not in export

    def test_it_is_the_forward_map_the_writer_uses(self, tmp_path):
        """Recomputed through the very function that names an export, so the
        two cannot disagree — and no fake is walked backwards."""
        case = _case(tmp_path, "Feit Decl. ISO Mot..docx")
        pz = _pz("Ondine Feit")
        named = P._pseudonymized_txt_path(
            case / "Text Files", case / "Feit Decl. ISO Mot..docx", pz, log)
        assert named.name.lower() in P._pn_export_sources(
            case, "Text Files", pz, log)

    def test_two_sources_scrubbing_alike_still_map_one_each(self, tmp_path):
        """The map cannot be ambiguous, and that is why it goes through
        `_pseudonymized_txt_path`: two stems that scrub alike are given
        different export names by the digest it appends."""
        case = _case(tmp_path)
        _docx(case, "Feit Decl..docx")
        _docx(case, "Feit_Decl..docx")        # same stem once separators go
        pz = _pz("Ondine Feit")
        m = P._pn_export_sources(case, "Text Files", pz, log)
        assert sorted(v for v in m.values() if v.startswith("Feit")) == [
            "Feit Decl..docx", "Feit_Decl..docx"]


class TestTheNameLookup:
    def test_an_export_and_its_quarantine_are_one_document(self):
        m = {"kingscote decl..txt": "Feit Decl..pdf"}
        for n in ("Kingscote Decl..txt", "Kingscote Decl..txt.LEAK"):
            assert P._pn_export_source_name(
                pathlib.Path("/x/Text Files") / n, m) == "Feit Decl..pdf"

    def test_an_export_of_no_source_keeps_its_own_name(self):
        """An orphan names itself — there is nothing truer to say."""
        assert P._pn_export_source_name(
            pathlib.Path("/x/Stray.txt"), {}) == "Stray.txt"


# ── both passes ─────────────────────────────────────────────────────────────

def _rows(case):
    ws = openpyxl.load_workbook(case / "LEAKS.xlsx").active
    head = [str(h).strip()
            for h in next(ws.iter_rows(max_row=1, values_only=True))]
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def _run(case, monkeypatch, tmp_path):
    monkeypatch.setenv("PDF_LINKER_MASTER", str(tmp_path / "master.xlsx"))
    monkeypatch.setattr(sys, "argv", ["pdf_linker.py", str(case)])
    try:
        P.main()
    except SystemExit:
        pass


class TestBothPassesNameTheSource:
    def test_the_full_run_does(self, tmp_path, monkeypatch):
        case = _case(tmp_path, "Feit Decl. ISO Mot..docx")
        _run(case, monkeypatch, tmp_path)
        assert [r["File"] for r in _rows(case)] == ["Feit Decl. ISO Mot..docx"]

    def test_and_so_does_apply_leak_fixes(self, tmp_path, monkeypatch):
        """The pass that produced the reported cells: a `yes` it refuses
        holds the file and rewrites the worksheet."""
        case = _case(tmp_path, "Feit Decl. ISO Mot..docx")
        _run(case, monkeypatch, tmp_path)
        wb = openpyxl.load_workbook(case / "LEAKS.xlsx")
        ws = wb.active
        head = [str(c.value).strip() for c in ws[1]]
        ws.cell(2, head.index("Fix? (yes/no)") + 1).value = "yes"
        wb.save(case / "LEAKS.xlsx")

        args = types.SimpleNamespace(term=[],
                                     key=str(case / "pseudonym_key.xlsx"))
        P._fix_leaks_mode(case, args, {}, log)
        files = [r["File"] for r in _rows(case)]
        assert files == ["Feit Decl. ISO Mot..docx"]
        assert not any("Kingscote" in f or f.endswith(".txt") for f in files)

    def test_the_export_itself_still_carries_no_real_name(self, tmp_path,
                                                          monkeypatch):
        """The worksheet is triage and stays in the case folder; the EXPORT is
        the shareable artifact and its own name is unchanged."""
        case = _case(tmp_path, "Feit Decl. ISO Mot..docx")
        _run(case, monkeypatch, tmp_path)
        names = [p.name for p in (case / "Text Files").glob("*.txt")]
        assert names and not any("Feit" in n for n in names)
