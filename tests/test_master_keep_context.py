"""The master KEEP sheet carries the KEPT TEXT and the SENTENCE it was reached
from — and nothing else of the case.

Two halves of one thing.

**The rows name the kept text.** A `no` or a `never` names the whole value and
that string IS the instruction. A keep-SPEC says two things at once:
`Alder Law, P.C. -> [Law]` records a lesson that generalises ("Law is never a
name") and a remainder that does not ("Alder" is this matter's law firm) — and
the remainder has never applied anywhere else, since an inherited keep-spec
builds no fragment terms. The sheet stored the whole value anyway, so the one
workbook that lives OUTSIDE every case folder, is routinely synced and is never
pruned carried that matter's party name. It is written under its kept parts
now, one row each, and a sheet an older version wrote HEALS on the next run.

**And each keep carries the sentence it was reached from.** A `no`/`never` —
typed into `LEAKS.xlsx`, or into the pseudonym key's Replacement column before
a re-run — is the operator saying the tool flagged or faked something it should
not have. The verdict alone cannot say why: "Charge" is boilerplate in "CHARGE
OF DISCRIMINATION" and a surname in "served on Charge at his residence". So the
sentence is kept beside it, on its own sheet, accumulated across matters, and
the review tiers can be tuned from real false positives. It is read from the
SCRUBBED EXPORT — every other real value in it already stands as its stand-in,
and the one real value left is the flagged one, which the operator has just
declared is not a name.

Run:  cd PDF-Linker && python3 -m pytest tests/test_master_keep_context.py -v
"""
import inspect
import logging

import openpyxl
import pytest

import pdf_linker as P

log = logging.getLogger("test")


# ── helpers ──────────────────────────────────────────────────────────────────

def _decision(value, cell, **kw):
    """One decision as `_pn_parse_decision_rows` would hand it back."""
    rows = [("Value", "Fix? (yes/no)", "Type", "Notes"),
            (value, cell, kw.pop("type", ""), kw.pop("notes", ""))]
    d = P._pn_parse_decision_rows(rows, log)[value.lower()]
    d.update(kw)
    return d


def _pz(names=("Rachel Ashworth",), keeps=()):
    reg = P._PnFakeRegistry()
    terms = P._pn_build_terms(list(names), [], [], registry=reg)
    pz = P.Pseudonymizer(terms, {}, registry=reg)
    decisions = {d["value"].lower(): d for d in keeps}
    pz._keep_decisions = {vl: d for vl, d in decisions.items()
                          if P._pn_decision_is_keep(d)}
    # Made HERE: the sample is evidence for the decision being made, and a
    # folder that merely inherits a keep flagged nothing to record.
    pz._keep_local = set(pz._keep_decisions)
    pz.keep_strict, pz.keep_soft, pz.keep_nuclear = P._pn_keep_values(decisions)
    return pz


BODY = ("====== Page 1 ======\n"
        " 1  The Charge of Discrimination was served on Rachel Ashworth at\n"
        " 2  her residence in Van Nuys on the fourth of July.\n")


def _sheet(path, title):
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        if ws.title.strip().lower() == title.strip().lower():
            return [list(r) for r in ws.iter_rows(values_only=True)]
    return []


def _keep_rows(path):
    return _sheet(path, P._PN_MASTER_KEEP_SHEET)[1:]


def _ctx_rows(path):
    return _sheet(path, P._PN_MASTER_KEEP_CONTEXT_SHEET)[1:]


# ── the rows name the kept text ──────────────────────────────────────────────

class TestTheRowNamesTheKeptText:
    def test_a_bracket_stores_the_fragment_and_not_the_party(self, tmp_path,
                                                             monkeypatch):
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        d = _decision("Alder Law, P.C.", "[Law]")
        P._pn_update_master_keep({}, {d["value"].lower(): d},
                                 "Case A", "2026-01-01", log)
        rows = _keep_rows(mp)
        assert [r[0] for r in rows] == ["Law"]
        assert rows[0][1] == "[Law]"                 # rebuilt from the part
        assert rows[0][2] == P._PN_KEEP_PART_TYPE
        assert not any("Alder" in str(c) for r in rows for c in r)

    def test_two_kept_parts_are_two_rows_with_their_own_promises(self, tmp_path,
                                                                 monkeypatch):
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        # A bracket and a brace in one cell are two different promises about
        # two different fragments, so they cannot share a row.
        d = _decision("Alder Law Group, P.C.", "[Law] {Group}")
        P._pn_update_master_keep({}, {d["value"].lower(): d},
                                 "Case A", "2026-01-01", log)
        got = {r[0]: (r[1], r[2]) for r in _keep_rows(mp)}
        assert got == {"Law": ("[Law]", P._PN_KEEP_PART_TYPE),
                       "Group": ("{Group}", P._PN_KEEP_NUCLEAR_TYPE)}

    def test_a_no_and_a_never_are_stored_whole(self, tmp_path, monkeypatch):
        """The whole value IS the instruction there — there is nothing to
        reduce, and a row without it could not be re-applied."""
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        ds = [_decision("Charge", "no"), _decision("Cross River Bank", "never")]
        P._pn_update_master_keep({}, {d["value"].lower(): d for d in ds},
                                 "Case A", "2026-01-01", log)
        got = {r[0]: r[1] for r in _keep_rows(mp)}
        assert got == {"Charge": "no", "Cross River Bank": "never"}

    def test_an_ocr_fix_keeps_its_garble(self, tmp_path, monkeypatch):
        """The one exception: there the VALUE is the garble the correction is
        defined against, so a row reduced to its kept part would no longer say
        what it corrects."""
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        d = _decision("avidsaid", "*David {said}")
        P._pn_update_master_keep({}, {d["value"].lower(): d},
                                 "Case A", "2026-01-01", log)
        assert [r[0] for r in _keep_rows(mp)] == ["avidsaid"]

    def test_an_alias_spec_takes_the_reduction(self, tmp_path, monkeypatch):
        """…where an ALIAS one does not need the exception: the alias half is
        case-local and is inherited by nobody, so only the brace survives."""
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        d = _decision("avidsaid", "~David {said}")
        P._pn_update_master_keep({}, {d["value"].lower(): d},
                                 "Case A", "2026-01-01", log)
        rows = _keep_rows(mp)
        assert [r[0] for r in rows] == ["said"]
        assert not any("David" in str(c) for r in rows for c in r)

    def test_an_older_sheet_heals(self, tmp_path, monkeypatch):
        """A workbook written before the rule sheds the remainder on the next
        run in ANY folder, carrying the accumulated history onto the part."""
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = P._PN_MASTER_KEEP_SHEET
        ws.append(list(P._PN_MASTER_KEEP_HEADERS))
        ws.append(["Alder Law, P.C.", "[Law]", "KEEP-PART", 4, "Case Z",
                   "2025-01-01", "2025-06-06", "cited authority: …", "Case Z"])
        wb.save(mp)
        P._pn_update_master_keep({}, {}, "Case A", "2026-01-01", log)
        rows = _keep_rows(mp)
        assert [r[0] for r in rows] == ["Law"]
        assert rows[0][3] == 4 and rows[0][4] == "Case Z"   # history carried
        assert rows[0][5] == "2025-01-01" and rows[0][6] == "2025-06-06"
        # The note described the value this row is no longer about.
        assert not rows[0][7]
        assert not any("Alder" in str(c) for r in rows for c in r)

    def test_a_note_naming_another_party_is_not_stored(self, tmp_path,
                                                        monkeypatch):
        """A pre-fill note names the tracked value a spelling was read as a
        misspelling of, and one survives onto a keep wherever the operator
        typed `no` over a pre-filled alias — so the cell explaining the row
        would carry the party name the row itself no longer does. The
        AUTHORITY note names a published decision, which is public record, and
        stays."""
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        note = (P._PN_PREFILL_NOTE.format(canon="Vazquez") + P._PN_NOTE_SEP
                + "cited authority: Kremerman v. White (2021) 71 Cal.App.5th 358")
        d = _decision("Vazqez", "no", notes=note)
        P._pn_update_master_keep({}, {"vazqez": d}, "Case A", "2026-01-01", log)
        (row,) = _keep_rows(mp)
        assert "Vazquez" not in str(row[7])
        assert "Kremerman" in str(row[7])

    def test_the_bracket_still_reads_back_as_a_strict_keep(self):
        """The tier has to survive the reduction: a whole-value bracket parses
        as a `no` because the cut left nothing over, and it is still a bracket
        — "this fragment is never a name", kept even beside a name."""
        d = _decision("Human Resources", "[Human Resources]")
        assert d["fix"] == "no"
        strict, soft, nuclear = P._pn_keep_values({"human resources": d})
        assert strict == {"Human Resources"}
        assert not soft and not nuclear

    def test_a_whole_value_brace_is_still_nuclear(self):
        d = _decision("Law", "{Law}")
        strict, soft, nuclear = P._pn_keep_values({"law": d})
        assert nuclear == {"Law"} and not soft


# ── the sentence each keep was reached from ──────────────────────────────────

class TestTheSample:
    def test_the_quote_is_the_export_and_not_the_document(self):
        pz = _pz(keeps=[_decision("Charge", "no", type="REVIEW name")])
        out = pz.apply(BODY)
        pz.note_keep_context(P._pn_body_lines(out))
        (sample,) = pz.keep_context["charge"]
        assert "Charge of Discrimination" in sample["quote"]
        # The party beside it stands as its stand-in, which is the whole of
        # what makes a permanent cross-case sample safe to hold.
        assert "Rachel Ashworth" not in sample["quote"]
        assert "Ashworth" not in sample["quote"]
        assert sample["keep"] == "no"
        assert sample["kind"] == "REVIEW name"    # the tier that flagged it

    def test_a_sentence_carrying_a_leaked_real_is_not_kept(self):
        """A sample is a substring of the export, so it is as shareable as the
        deliverable — except that the file may be about to be quarantined for
        a leak. Asked of the SENTENCE, so a leak three pages away costs
        nothing and a leak in this one drops the sample."""
        pz = _pz(keeps=[_decision("Charge", "no")])
        assert pz._keep_sample_is_clean("The Charge was served on Fenmore.")
        assert not pz._keep_sample_is_clean(
            "The Charge was served on Rachel Ashworth.")

    def test_an_inherited_keep_is_not_sampled(self):
        """A folder that merely inherits a keep flagged nothing, so there is no
        false positive there to record — and a master sheet carries hundreds of
        inherited keeps, several of them ordinary vocabulary."""
        pz = _pz(keeps=[_decision("Charge", "no")])
        pz._keep_local = set()                      # inherited, not typed here
        pz.note_keep_context(P._pn_body_lines(pz.apply(BODY)))
        assert not pz.keep_context

    def test_a_keep_that_protected_nothing_here_yields_nothing(self):
        pz = _pz(keeps=[_decision("Registry", "no")])
        pz.note_keep_context(P._pn_body_lines(pz.apply(BODY)))
        assert not pz.keep_context.get("registry")

    def test_a_keep_spec_is_sampled_under_its_kept_part(self):
        pz = _pz(names=("Alder Law, P.C.",),
                 keeps=[_decision("Alder Law, P.C.", "[Law]")])
        body = ("====== Page 1 ======\n"
                " 1  The retainer was signed by Alder Law, P.C. in March.\n")
        pz.note_keep_context(P._pn_body_lines(pz.apply(body)))
        assert "alder law, p.c." not in pz.keep_context
        (sample,) = pz.keep_context["Law".lower()]
        assert "Alder" not in sample["quote"]
        assert "Law" in sample["quote"]

    def test_one_sentence_is_kept_once_however_often_it_recurs(self):
        pz = _pz(keeps=[_decision("Charge", "no")])
        out = pz.apply(BODY)
        for _ in range(3):
            pz.note_keep_context(P._pn_body_lines(out))
        assert len(pz.keep_context["charge"]) == 1

    def test_a_value_is_sampled_at_most_the_cap(self):
        pz = _pz(keeps=[_decision("Charge", "no")])
        for n in range(P._PN_KEEP_CONTEXT_MAX + 3):
            body = (f"====== Page 1 ======\n"
                    f" 1  Exhibit {n} records that the Charge was filed in "
                    f"the {n}th week of the year.\n")
            pz.note_keep_context(P._pn_body_lines(pz.apply(body)))
        assert len(pz.keep_context["charge"]) == P._PN_KEEP_CONTEXT_MAX


# ── …and it lands on its own sheet ───────────────────────────────────────────

class TestTheSheet:
    def _write(self, mp, samples, case="Case A", today="2026-01-01"):
        d = _decision("Charge", "no")
        P._pn_update_master_keep({}, {"charge": d}, case, today, log,
                                 contexts=samples)

    def test_the_sample_lands_beside_the_decision(self, tmp_path, monkeypatch):
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        self._write(mp, [{"value": "Charge", "keep": "no",
                          "kind": "REVIEW name",
                          "quote": "The Charge of Discrimination was served."}])
        rows = _ctx_rows(mp)
        assert rows == [["Charge", "no", "REVIEW name",
                         "The Charge of Discrimination was served.",
                         "Case A", "2026-01-01"]]
        # …without disturbing the decision sheet beside it.
        assert [r[0] for r in _keep_rows(mp)] == ["Charge"]

    def test_a_rerun_of_one_folder_does_not_add_the_sentence_twice(
            self, tmp_path, monkeypatch):
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        s = [{"value": "Charge", "keep": "no", "kind": "",
              "quote": "The Charge of Discrimination was served."}]
        self._write(mp, s, today="2026-01-01")
        self._write(mp, s, today="2026-02-02")
        assert len(_ctx_rows(mp)) == 1

    def test_another_matter_adds_its_own_sentence(self, tmp_path, monkeypatch):
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        self._write(mp, [{"value": "Charge", "keep": "no", "kind": "",
                          "quote": "The Charge of Discrimination was served."}],
                    case="Case A")
        self._write(mp, [{"value": "Charge", "keep": "no", "kind": "",
                          "quote": "Served on Charge at his residence."}],
                    case="Case B")
        rows = _ctx_rows(mp)
        assert len(rows) == 2
        assert {r[4] for r in rows} == {"Case A", "Case B"}

    def test_the_sheet_is_capped_across_runs(self, tmp_path, monkeypatch):
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        for n in range(P._PN_KEEP_CONTEXT_MAX + 4):
            self._write(mp, [{"value": "Charge", "keep": "no", "kind": "",
                              "quote": f"Sentence number {n} about the Charge."}],
                        case=f"Case {n}")
        assert len(_ctx_rows(mp)) == P._PN_KEEP_CONTEXT_MAX

    def test_no_samples_no_sheet(self, tmp_path, monkeypatch):
        mp = tmp_path / "master_leaks.xlsx"
        monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
        self._write(mp, [])
        assert not _sheet(mp, P._PN_MASTER_KEEP_CONTEXT_SHEET)


# ── …from every writer, pinned on the source ─────────────────────────────────

def test_every_export_writer_records_the_samples():
    """Three writers build the same leak evidence from the same scrubbed body,
    and a pass missing from one of two paths is this project's oldest shape of
    bug (the Word path ran every scan and none of the cures). Pinned here so
    the next writer cannot be added without it."""
    for fn in (P._write_text_version, P._write_word_text_version,
               P._fix_leaks_mode):
        src = inspect.getsource(fn)
        assert ".note_keep_context(" in src, (
            f"{fn.__name__} does not record its KEEP context samples")


def test_the_samples_reach_the_master_workbook_from_both_run_sites():
    for fn in (P._fix_leaks_mode, P.main):
        src = inspect.getsource(fn)
        assert "contexts=_pn_keep_context_samples(" in src, (
            f"{fn.__name__} writes the KEEP sheet without its samples")


# ── end to end, through the pass the operator clicks ─────────────────────────

def _args(folder):
    import argparse
    return argparse.Namespace(folder=str(folder), fix_leaks=True, term=[],
                              key=None, key_out=None, no_pseudonymize=False,
                              detectors=None, dump_terms=False, first=False)


def test_a_no_typed_in_the_worksheet_lands_with_its_sentence(tmp_path,
                                                             monkeypatch):
    """The operator's own path: a value flagged in `LEAKS.xlsx`, answered `no`,
    Apply Leak Fixes clicked. The verdict reaches the KEEP sheet and the
    sentence it was reached from reaches the one beside it — with the party in
    that sentence standing as its stand-in, which is what makes the row safe to
    keep in a permanent cross-case workbook."""
    mp = tmp_path / "master_leaks.xlsx"
    monkeypatch.setenv("PDF_LINKER_MASTER", str(mp))
    folder = tmp_path / "Case"
    tdir = folder / "Text Files"
    tdir.mkdir(parents=True)
    (tdir / "Brief.txt.LEAK").write_text(
        "====== Page 1 ======\n"
        " 1  The Charge of Discrimination was served on Rachel Ashworth at\n"
        " 2  her residence in Van Nuys on the fourth of July.\n",
        encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pseudonym Key"
    ws.append(["Category", "Real Value", "Replacement", "Status", "Source",
               "Occurrences"])
    ws.append(["person", "Rachel Ashworth", "Chilcott Weatherby", "replaced",
               "spreadsheet", 1])
    wb.save(folder / "pseudonym_key.xlsx")
    wb2 = openpyxl.Workbook()
    w2 = wb2.active
    w2.title = "Potential Leaks"
    w2.append(["File", "Type", "Value", "Where (page:line)",
               "Fix? (yes/no)", "Notes"])
    w2.append(["Brief.txt.LEAK", "REVIEW name", "Charge", "p.1:1", "no", ""])
    wb2.save(folder / "LEAKS.xlsx")

    P._fix_leaks_mode(folder, _args(folder), {}, log)

    assert [r[0] for r in _keep_rows(mp)] == ["Charge"]
    (row,) = _ctx_rows(mp)
    assert row[0] == "Charge" and row[1] == "no"
    assert row[2] == "REVIEW name"                 # the tier that flagged it
    assert "Charge of Discrimination" in row[3]
    assert "Rachel Ashworth" not in row[3]         # the party stands as a fake
    assert "Chilcott Weatherby" in row[3]
