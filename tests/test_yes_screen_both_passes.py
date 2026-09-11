"""A worksheet `yes` is screened the SAME WAY by both passes.

`_pn_vocabulary_screen` shipped in `--fix-leaks` alone, so a value the
text-only pass refused as vocabulary — "NAN", "JTII" — was minted as a person
by the FULL run on the very next click. The operator's two buttons disagreed
about one cell, and the disagreement was silent on the path that acted.

One helper decides it now (`_pn_yes_refusal`) and both passes ask it, the way
`Pseudonymizer.leak_findings` is the one list of REVIEW tiers. The direction is
the safe one: a wrong refusal costs one worksheet row the operator answers
again with a replacement, where a wrong mint rewrites the document's own
vocabulary as a surname in every export.

Run:  cd PDF-Linker && python3 -m pytest tests/test_yes_screen_both_passes.py -v
"""
import inspect
import logging
import os
import re
import sys
import types
import zipfile

import openpyxl

import pdf_linker as P

log = logging.getLogger("test")
LEAK_HDR = ("Value", "Fix? (yes/no)", "File", "Type", "Where (page:line)",
            "Notes")
SCREEN = "it is a lone all-caps token of four letters or fewer"


class _Log:
    def __init__(self):
        self.lines = []

    def warning(self, msg):
        self.lines.append(msg)

    info = debug = warning


# ── one decision site ───────────────────────────────────────────────────────

class TestTheScreenIsAskedInOnePlace:
    def test_both_passes_call_the_helper(self):
        src = inspect.getsource(P)
        assert "_pn_yes_refusal(d, orig_vocab, log, \"--fix-leaks\")" in src
        assert "_pn_yes_refusal(d, orig_vocab, log, \"Leak worksheet\")" in src

    def test_the_screen_itself_is_consulted_nowhere_else(self):
        """A second caller of the screen is the two paths splitting again."""
        module = inspect.getsource(P)
        # Built once per pass; asked only inside the helper.
        assert len(re.findall(r"_pn_vocabulary_screen\(", module)) == 3   # def + 2 builds
        body = inspect.getsource(P._pn_yes_refusal)
        assert "vocab_screen(d[\"value\"])" in body
        assert len(re.findall(r"orig_vocab\(", module)) == 0

    def test_the_reason_is_the_rows_own_wording(self):
        """One string, so the log and LEAKS.xlsx cannot say different things."""
        assert "_PN_REFUSED_VOCAB.format(" in inspect.getsource(P._pn_yes_refusal)


# ── the helper ──────────────────────────────────────────────────────────────

class TestTheHelper:
    def test_it_refuses_and_says_why_and_how_to_proceed(self):
        lg = _Log()
        why = P._pn_yes_refusal({"value": "NAN"}, P._pn_vocabulary_screen([]),
                                lg, "--fix-leaks")
        assert SCREEN in why
        assert "~CANONICAL" in why
        assert lg.lines and SCREEN in lg.lines[0]

    def test_a_value_the_screen_allows_is_minted_silently(self):
        lg = _Log()
        assert P._pn_yes_refusal({"value": "Rouzbahni"},
                                 P._pn_vocabulary_screen([]), lg, "x") == ""
        assert lg.lines == []

    def test_a_phrase_is_never_screened_as_a_blanket_yes(self):
        """`phrase` is the operator's deliberate statement about a name."""
        lg = _Log()
        assert P._pn_yes_refusal({"value": "NAN", "phrase": True},
                                 P._pn_vocabulary_screen([]), lg, "x") == ""
        assert lg.lines == []


# ── end to end, through the FULL run ────────────────────────────────────────

def _word_case(case, body_text):
    case.mkdir()
    body = "".join(f"<w:p><w:r><w:t>{t}</w:t></w:r></w:p>"
                   for t in body_text.split("\n"))
    with zipfile.ZipFile(case / "Motion.docx", "w") as z:
        z.writestr("word/document.xml",
                   '<?xml version="1.0"?><w:document xmlns:w="http://schemas.'
                   'openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                   + body + "</w:body></w:document>")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Case Number", "Title Plaintiff", "Title Defendant"])
    ws.append(["24STCV00001", "Hollis Vantreight", "Cascadia Freight, Inc."])
    wb.save(case / "Order_Mine.xlsx")
    return case


BODY = ("Plaintiff Hollis Vantreight sued Cascadia Freight, Inc.\n"
        "Respondent materially breached the Agreement as NAN alleged supra.\n")


def _leaks(case, *rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LEAKS"
    ws.append(list(LEAK_HDR))
    for value, cell in rows:
        ws.append([value, cell, "Motion.docx", "narrative name?", "line 2", ""])
    wb.save(case / "LEAKS.xlsx")


def _run(case, monkeypatch, tmp_path):
    monkeypatch.setenv("PDF_LINKER_MASTER", str(tmp_path / "master.xlsx"))
    monkeypatch.setattr(sys, "argv", ["pdf_linker.py", str(case)])
    try:
        P.main()
    except SystemExit:
        pass


def _export(case):
    return (case / "Text Files" / "Motion.txt").read_text(encoding="utf-8")


def _row(case, value):
    ws = openpyxl.load_workbook(case / "LEAKS.xlsx").active
    head = [str(h).strip()
            for h in next(ws.iter_rows(max_row=1, values_only=True))]
    v = head.index("Value")
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[v] == value:
            return dict(zip(head, r))
    raise AssertionError(f"no row for {value!r}")


class TestTheFullRunScreensToo:
    def test_it_does_not_mint_the_refused_value(self, tmp_path, monkeypatch):
        case = _word_case(tmp_path / "Case A", BODY)
        _leaks(case, ("NAN", "yes"))
        _run(case, monkeypatch, tmp_path)
        out = _export(case)
        assert "NAN" in out                       # not minted as a person
        assert "Hollis" not in out                # the real party still is

    def test_and_the_row_comes_back_saying_why(self, tmp_path, monkeypatch):
        case = _word_case(tmp_path / "Case A", BODY)
        _leaks(case, ("NAN", "yes"))
        _run(case, monkeypatch, tmp_path)
        notes = str(_row(case, "NAN")["Notes"])
        assert notes.startswith(P._PN_REFUSED_NOTE_LEAD)
        assert SCREEN in notes

    def test_a_value_the_screen_allows_is_still_scrubbed(self, tmp_path,
                                                         monkeypatch):
        """The screen refuses a class, not the feature: an ordinary `yes`
        mints as it always did."""
        case = _word_case(tmp_path / "Case B", BODY.replace(
            "NAN", "Rouzbahni"))
        _leaks(case, ("Rouzbahni", "yes"))
        _run(case, monkeypatch, tmp_path)
        assert "Rouzbahni" not in _export(case)


class TestBothPassesRefuseTheSameValue:
    def test_the_two_agree(self, tmp_path, monkeypatch):
        """The whole point: what one pass refuses, the other refuses too."""
        case = _word_case(tmp_path / "Case A", BODY)
        _leaks(case, ("NAN", "yes"))
        _run(case, monkeypatch, tmp_path)
        assert "NAN" in _export(case)

        # …and Apply Leak Fixes, on the folder the full run just left.
        args = types.SimpleNamespace(
            term=[], key=str(case / "pseudonym_key.xlsx"))
        monkeypatch.setenv("PDF_LINKER_MASTER", str(tmp_path / "master.xlsx"))
        P._fix_leaks_mode(case, args, {}, log)
        assert "NAN" in _export(case)
        assert SCREEN in str(_row(case, "NAN")["Notes"])
